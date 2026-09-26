"""Unit and mocked-route tests for LectureMind. No models or network are used."""

import io
import os
import sys
import unittest
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from pptx import Presentation

import app as pipeline
import webapp


SUMMARY = """KEY CONCEPTS
1. Recursion: A function can call itself to solve a smaller case.
2. Base case: A stopping condition prevents endless recursion.
SUMMARY
The lecture explains recursive problem solving and the role of a base case."""

FLASHCARDS = """Front: What is recursion?
Back: A function calling itself to solve a smaller case.

Front: What does a base case do?
Back: It stops the recursion."""

QUIZ = """Question: What prevents a recursive function from continuing forever?
A. A variable
B. A base case
C. A loop counter
D. A comment
Answer: B"""


class PipelineUnitTests(unittest.TestCase):
    def test_blip_loader_uses_processor_and_caption_model(self):
        processor, model = object(), object()
        processor_loader = Mock(return_value=processor)
        model_loader = Mock(return_value=model)
        fake_transformers = SimpleNamespace(
            BlipProcessor=SimpleNamespace(from_pretrained=processor_loader),
            BlipForConditionalGeneration=SimpleNamespace(from_pretrained=model_loader),
        )

        with patch.dict(sys.modules, {"transformers": fake_transformers}):
            with patch.object(pipeline, "_captioner", None), patch.object(pipeline, "_captioner_attempted", False):
                loaded = pipeline._load_captioner()

        self.assertEqual(loaded, (processor, model))
        processor_loader.assert_called_once_with("Salesforce/blip-image-captioning-base")
        model_loader.assert_called_once_with("Salesforce/blip-image-captioning-base")

    def test_split_chunks_preserves_words_and_respects_limit(self):
        source = "word " * 100
        chunks = pipeline._split_chunks(source, limit=37)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk) <= 37 for chunk in chunks))
        self.assertEqual(" ".join(chunks), source.strip())

    @patch.object(pipeline, "_chat", return_value="Condensed lecture section.")
    def test_build_context_summarizes_long_material_in_chunks(self, chat):
        context = pipeline.build_context(transcript="Important lecture fact. " * 500)

        self.assertIn("hierarchical summaries", context)
        self.assertIn("Condensed lecture section", context)
        self.assertGreater(chat.call_count, 1)

    def test_parse_overview_separates_concepts_from_summary_and_removes_markdown(self):
        parsed = pipeline.parse_overview(
            "KEY CONCEPTS\n1. **Recursion**: A function calls itself.\n"
            "2. Base case: A condition stops recursion.\nSUMMARY\n"
            "The lecture introduces recursion and explains how the base case stops it."
        )

        self.assertEqual([item["label"] for item in parsed["concepts"]], ["Recursion", "Base case"])
        self.assertEqual(parsed["concepts"][0]["detail"], "A function calls itself.")
        self.assertEqual(parsed["summary"], "The lecture introduces recursion and explains how the base case stops it.")

    def test_parse_flashcards_accepts_markdown_and_bullets(self):
        parsed = pipeline.parse_flashcards("- **Front:** What is a base case?\n- **Back:** A stopping condition.")

        self.assertEqual(parsed, [{"front": "What is a base case?", "back": "A stopping condition."}])

    def test_parse_quiz_extracts_question_options_and_answer(self):
        parsed = pipeline.parse_quiz(QUIZ)

        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0]["prompt"], "What prevents a recursive function from continuing forever?")
        self.assertEqual(parsed[0]["options"][1], {"letter": "B", "text": "A base case"})
        self.assertEqual(parsed[0]["answer"], "B")

    def test_mindmap_concepts_keep_labels_separate_from_explanations(self):
        concepts = pipeline._concepts_from_summary(
            "1. **Recursion**: A function calls itself.\n2. Base case: A stopping condition."
        )

        self.assertEqual(concepts[0], {"label": "Recursion", "details": ["A function calls itself."]})
        self.assertEqual(concepts[1]["label"], "Base case")

    def test_fallback_mindmap_escapes_labels_and_places_nodes_on_both_sides(self):
        svg = pipeline._fallback_svg("Computer Science & Design", [
            {"label": "History <and> design", "details": ["Early ideas"]},
            {"label": "Operating systems", "details": ["Manage computer resources"]},
            {"label": "Networks", "details": ["Connect devices"]},
        ])

        self.assertIn("Computer Science &amp;</text>", svg)
        self.assertIn("History &lt;and&gt; design", svg)
        self.assertIn('x="30"', svg)
        self.assertIn('x="900"', svg)
        self.assertIn("</svg>", svg)


class FlaskRouteTests(unittest.TestCase):
    def setUp(self):
        webapp.app.config.update(TESTING=True, MAX_CONTENT_LENGTH=500 * 1024 * 1024)
        self.client = webapp.app.test_client()

    @staticmethod
    def _valid_pptx():
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        stream = io.BytesIO()
        presentation.save(stream)
        stream.seek(0)
        return stream

    def _mock_generation(self, stack):
        stack.enter_context(patch.object(webapp, "build_context", return_value="Mock lecture context"))
        stack.enter_context(patch.object(webapp, "generate_summary", return_value=SUMMARY))
        stack.enter_context(patch.object(webapp, "generate_flashcards", return_value=FLASHCARDS))
        stack.enter_context(patch.object(webapp, "generate_quiz", return_value=QUIZ))
        stack.enter_context(patch.object(webapp, "generate_mindmap", return_value={
            "title": "Recursion", "concepts": [{"label": "Base case", "details": ["Stops recursion"]}]
        }))
        stack.enter_context(patch.object(webapp, "render_mindmap_svg", return_value="<svg></svg>"))

    def test_home_page_loads(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"LectureMind", response.data)

    def test_process_rejects_empty_and_unsupported_uploads(self):
        empty = self.client.post("/process", data={})
        unsupported = self.client.post("/process", data={
            "audio": (io.BytesIO(b"text"), "notes.txt")
        }, content_type="multipart/form-data")

        self.assertEqual(empty.status_code, 400)
        self.assertIn(b"Choose an audio recording", empty.data)
        self.assertEqual(unsupported.status_code, 400)
        self.assertIn(b".mp3", unsupported.data)

    def test_audio_workflow_cleans_uploaded_file_after_request(self):
        saved_paths = []

        def fake_transcribe(path):
            saved_paths.append(Path(path))
            return "A transcript about recursion."

        with ExitStack() as stack:
            self._mock_generation(stack)
            stack.enter_context(patch.object(webapp, "transcribe", side_effect=fake_transcribe))
            response = self.client.post("/process", data={
                "audio": (io.BytesIO(b"fake audio bytes"), "lecture.wav")
            }, content_type="multipart/form-data")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Knowledge check", response.data)
        self.assertIn(b"Key concepts", response.data)
        self.assertIn(b"Show answer", response.data)
        self.assertIn(b'type="radio"', response.data)
        self.assertIn(b'name="quiz-question-1"', response.data)
        self.assertIn(b"Check answer", response.data)
        self.assertIn(b"reveal it when", response.data)
        self.assertEqual(len(saved_paths), 1)
        self.assertFalse(saved_paths[0].exists())
        self.assertFalse(saved_paths[0].parent.exists())

    def test_slides_only_workflow(self):
        with ExitStack() as stack:
            self._mock_generation(stack)
            slide_processor = stack.enter_context(patch.object(webapp, "process_slides", return_value="Slide 1: recursion"))
            response = self.client.post("/process", data={
                "slides": (self._valid_pptx(), "lecture.pptx")
            }, content_type="multipart/form-data")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Slide notes", response.data)
        slide_processor.assert_called_once()

    def test_combined_workflow(self):
        with ExitStack() as stack:
            self._mock_generation(stack)
            transcriber = stack.enter_context(patch.object(webapp, "transcribe", return_value="Audio transcript"))
            slide_processor = stack.enter_context(patch.object(webapp, "process_slides", return_value="Slide text"))
            response = self.client.post("/process", data={
                "audio": (io.BytesIO(b"fake audio bytes"), "lecture.m4a"),
                "slides": (self._valid_pptx(), "lecture.pptx"),
            }, content_type="multipart/form-data")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Transcript", response.data)
        self.assertIn(b"Slide notes", response.data)
        transcriber.assert_called_once()
        slide_processor.assert_called_once()

    def test_damaged_powerpoint_is_rejected(self):
        response = self.client.post("/process", data={
            "slides": (io.BytesIO(b"not a powerpoint"), "broken.pptx")
        }, content_type="multipart/form-data")

        self.assertEqual(response.status_code, 400)
        self.assertIn(b"empty or damaged", response.data)

    def test_oversized_request_returns_clear_error(self):
        previous_limit = webapp.app.config["MAX_CONTENT_LENGTH"]
        webapp.app.config["MAX_CONTENT_LENGTH"] = 256
        try:
            response = self.client.post("/process", data={
                "audio": (io.BytesIO(b"a" * 1024), "lecture.wav")
            }, content_type="multipart/form-data")
        finally:
            webapp.app.config["MAX_CONTENT_LENGTH"] = previous_limit

        self.assertEqual(response.status_code, 413)
        self.assertIn(b"500 MB upload limit", response.data)

    def test_generation_failure_keeps_other_results(self):
        with ExitStack() as stack:
            self._mock_generation(stack)
            stack.enter_context(patch.object(webapp, "transcribe", return_value="Transcript remains available"))
            stack.enter_context(patch.object(webapp, "generate_flashcards", side_effect=pipeline.PipelineError("Flashcards could not be generated.")))
            response = self.client.post("/process", data={
                "audio": (io.BytesIO(b"fake audio bytes"), "lecture.mp3")
            }, content_type="multipart/form-data")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Knowledge check", response.data)
        self.assertIn(b"Flashcards could not be generated", response.data)
        self.assertIn(b"Transcript remains available", response.data)


if __name__ == "__main__":
    unittest.main()
